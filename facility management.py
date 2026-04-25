import streamlit as st
import pandas as pd
import json
from datetime import datetime
import os
import io
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH


# --- WORD REPORT GENERATOR ---
def generate_word_report(facility, report_date, areas):
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1.2)
        section.right_margin = Inches(1.2)

    # Title
    title = doc.add_heading("", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Facility and Security Report")
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(0x1A, 0x56, 0x8C)

    # Subtitle
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run(f"{facility}  |  Report Date: {report_date}").font.size = Pt(11)
    doc.add_paragraph()

    def add_section_heading(text):
        h = doc.add_heading(text, level=2)
        h.runs[0].font.color.rgb = RGBColor(0x1A, 0x56, 0x8C)

    def add_field(label, value):
        p = doc.add_paragraph()
        lr = p.add_run(f"{label}:  ")
        lr.bold = True
        lr.font.size = Pt(11)
        lr.font.color.rgb = RGBColor(0x1A, 0x56, 0x8C)
        vr = p.add_run(str(value) if value else "-")
        vr.font.size = Pt(11)

    def parse_personnel(raw):
        records = []
        if not raw or raw == "-":
            return records
        for entry in raw.split("; "):
            entry = entry.strip()
            if " (KES " in entry:
                name, rest = entry.rsplit(" (KES ", 1)
                amount = rest.rstrip(")")
                records.append((name.strip(), f"KES {amount}"))
        return records

    for idx, area in enumerate(areas):
        add_section_heading(f"Area of Concentration {idx+1}: {area['name']}")

        add_field("Current Status", area["status"])
        add_field("Improvements Needed", area["improvements"])

        # Budget table
        p = doc.add_paragraph()
        r = p.add_run("Quarterly Budget Allocation")
        r.bold = True
        r.font.size = Pt(11)
        tbl = doc.add_table(rows=2, cols=4)
        tbl.style = "Table Grid"
        headers = ["Q1 Budget", "Q2 Budget", "Q3 Budget", "Q4 Budget"]
        values = [f"KES {area['q1']:,}", f"KES {area['q2']:,}", f"KES {area['q3']:,}", f"KES {area['q4']:,}"]
        for i, (h, v) in enumerate(zip(headers, values)):
            cell = tbl.rows[0].cells[i]
            cell.text = h
            cell.paragraphs[0].runs[0].bold = True
            tbl.rows[1].cells[i].text = v
        doc.add_paragraph()

        # Personnel
        p = doc.add_paragraph()
        r = p.add_run("Personnel")
        r.bold = True
        r.font.size = Pt(11)

        personnel_str = area.get("personnel", "")
        if "EMPLOYED:" in personnel_str and "CASUALS:" in personnel_str:
            emp_part = personnel_str.split("EMPLOYED:")[1].split("|")[0].strip()
            cas_part = personnel_str.split("CASUALS:")[1].strip()
        else:
            emp_part, cas_part = "-", "-"

        for label, part in [("Employed Staff", emp_part), ("Casual Workers", cas_part)]:
            p2 = doc.add_paragraph()
            r2 = p2.add_run(label)
            r2.bold = True
            r2.font.size = Pt(10)
            records = parse_personnel(part)
            if records:
                ptbl = doc.add_table(rows=len(records)+1, cols=2)
                ptbl.style = "Table Grid"
                ptbl.rows[0].cells[0].text = "Name"
                ptbl.rows[0].cells[1].text = "Amount Paid"
                ptbl.rows[0].cells[0].paragraphs[0].runs[0].bold = True
                ptbl.rows[0].cells[1].paragraphs[0].runs[0].bold = True
                for ri, (name, amount) in enumerate(records):
                    ptbl.rows[ri+1].cells[0].text = name
                    ptbl.rows[ri+1].cells[1].text = amount
            else:
                doc.add_paragraph("No records.")
            doc.add_paragraph()

        add_field("Remarks", area["remarks"])
        doc.add_paragraph()
        doc.add_paragraph("─" * 60)
        doc.add_paragraph()

    # Footer
    footer_p = doc.add_paragraph()
    footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fr = footer_p.add_run(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    fr.font.size = Pt(9)
    fr.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


# --- 1. SYSTEM CONFIGURATION ---
st.set_page_config(page_title="Facility and Security System", layout="wide")

# --- RESPONSIVE CSS ---
st.markdown("""
<style>
    @media (max-width: 768px) {
        [data-testid="column"] {
            width: 100% !important;
            flex: 1 1 100% !important;
            min-width: 100% !important;
        }
        .stButton > button {
            width: 100% !important;
            padding: 0.75rem !important;
            font-size: 1rem !important;
        }
        .stTextInput input, .stNumberInput input, .stTextArea textarea {
            font-size: 1rem !important;
        }
        .block-container {
            padding-left: 1rem !important;
            padding-right: 1rem !important;
        }
    }
    .stButton > button { border-radius: 8px; }
    .area-box {
        border: 1px solid #1A568C;
        border-radius: 10px;
        padding: 1rem;
        margin-bottom: 1.5rem;
        background: #f7faff;
    }
</style>
""", unsafe_allow_html=True)

# --- FACILITIES ---
FACILITIES = ["KAG HQ", "KENYA KIDS", "KK TOWERS", "BURUBURU", "SECURITY"]

# --- 2. HEADER & NAVIGATION ---
st.title("🏗️ Facility and Security Report")
st.markdown("This system captures manual weekly data to feed the **CrewAI** backend.")

with st.sidebar:
    st.header("Site Navigation")
    current_site = st.radio("Select Facility to Update:", FACILITIES)
    report_date = st.date_input("Reporting Date", value=datetime.now())
    st.divider()
    st.success(f"Currently Logging: **{current_site}**")

    if st.checkbox("Show Recent Logs"):
        if os.path.exists("facility_master_log.csv"):
            recent_df = pd.read_csv("facility_master_log.csv")
            st.dataframe(recent_df.tail(5))

    st.divider()
    st.subheader("📥 Download Full Log")
    if os.path.exists("facility_master_log.csv"):
        full_df = pd.read_csv("facility_master_log.csv")
        csv_buffer = full_df.to_csv(index=False).encode("utf-8")
        st.download_button(
            label="⬇️ Download All Reports (CSV)",
            data=csv_buffer,
            file_name="facility_master_log.csv",
            mime="text/csv",
            use_container_width=True
        )
    else:
        st.info("No log file yet. Submit a report first.")

# --- 3. MAIN FORM ---
st.header(f"Weekly Update: {current_site}")

# Security facility: pre-load the 10 guideline areas
SECURITY_AREAS = [
    "1. Facilities Status (Site Overview)",
    "2. Maintenance and Repairs",
    "3. Projects and Infrastructure Improvements",
    "4. Security Management",
    "5. Security Incidents and Risk Monitoring",
    "6. Compliance and Safety Measures",
    "7. Vendors and Contractors",
    "8. Costs and Requisitions",
    "9. Challenges and Observations",
    "10. Next Steps / Priorities",
]

# Manage number of areas of concentration
if f"num_areas_{current_site}" not in st.session_state:
    if current_site == "SECURITY":
        st.session_state[f"num_areas_{current_site}"] = len(SECURITY_AREAS)
    else:
        st.session_state[f"num_areas_{current_site}"] = 1

col_add, col_remove = st.columns([1, 1])
with col_add:
    if st.button("➕ Add Area of Concentration", use_container_width=True):
        st.session_state[f"num_areas_{current_site}"] += 1
with col_remove:
    if st.button("➖ Remove Last Area", use_container_width=True):
        if st.session_state[f"num_areas_{current_site}"] > 1:
            st.session_state[f"num_areas_{current_site}"] -= 1

st.divider()

num_areas = st.session_state[f"num_areas_{current_site}"]
areas_data = []

for a in range(num_areas):
    st.markdown(f"### 📌 Area of Concentration {a+1}")

    default_area_name = SECURITY_AREAS[a] if current_site == "SECURITY" and a < len(SECURITY_AREAS) else ""
    area_name = st.text_input(
        "Field / Area Name",
        value=default_area_name,
        key=f"area_name_{a}_{current_site}",
        placeholder="e.g. Security, Plumbing, Electrical, Cleaning..."
    )

    SECURITY_STATUS_HINTS = [
        "Condition of facilities, areas in good condition vs needing attention, changes this week",
        "Number of tasks completed, type of works done, locations, status (completed/ongoing/pending)",
        "Ongoing projects, stage, timelines, progress, delays or risks",
        "Number of personnel on duty, deployment across sites, changes in security arrangements, status of security systems",
        "Incidents that occurred, actions taken, preventive measures introduced, emerging risks",
        "Safety measures implemented, compliance with standards, inspections carried out",
        "Contractors/suppliers engaged, nature of work done, performance issues",
        "Requisitions raised, estimated or actual costs, within budget, cost concerns",
        "Nature of issue, location, cause, suggested solution",
        "Clear, measurable next steps focused on improving facilities and security",
    ]
    SECURITY_IMPROVEMENT_HINTS = [
        "Any changes or improvements needed across sites",
        "Any major repairs or upgrades required",
        "Any risks or delays to address",
        "Any upgrades needed for gates, access control, perimeter, lighting, etc.",
        "Additional preventive measures to introduce",
        "Any gaps in compliance or safety to address",
        "Any contractor performance issues to resolve",
        "Any cost variances or budget concerns to address",
        "Suggested solutions for each challenge",
        "Priority actions for the coming week",
    ]
    status_hint = SECURITY_STATUS_HINTS[a] if current_site == "SECURITY" and a < len(SECURITY_STATUS_HINTS) else "Describe the current status of this area..."
    improvement_hint = SECURITY_IMPROVEMENT_HINTS[a] if current_site == "SECURITY" and a < len(SECURITY_IMPROVEMENT_HINTS) else "Detail repairs, upgrades, or maintenance needed"

    c1, c2 = st.columns(2)
    with c1:
        area_status = st.text_area(
            "Current Status",
            key=f"status_{a}_{current_site}",
            placeholder=status_hint
        )
    with c2:
        area_improvements = st.text_area(
            "Improvements Needed",
            key=f"improvements_{a}_{current_site}",
            placeholder=improvement_hint
        )

    st.markdown("**Quarterly Budget Allocation (KES)**")
    b1, b2, b3, b4 = st.columns(4)
    with b1:
        q1 = st.number_input("Q1", min_value=0, step=100, key=f"q1_{a}_{current_site}")
    with b2:
        q2 = st.number_input("Q2", min_value=0, step=100, key=f"q2_{a}_{current_site}")
    with b3:
        q3 = st.number_input("Q3", min_value=0, step=100, key=f"q3_{a}_{current_site}")
    with b4:
        q4 = st.number_input("Q4", min_value=0, step=100, key=f"q4_{a}_{current_site}")

    st.markdown("**Personnel**")
    p1, p2 = st.columns(2)

    with p1:
        st.markdown("👔 **Employed Staff**")
        num_employed = st.number_input("Number of Employed", min_value=0, max_value=20, step=1, key=f"num_emp_{a}_{current_site}")
        employed_records = []
        for i in range(int(num_employed)):
            ec1, ec2 = st.columns(2)
            with ec1:
                emp_name = st.text_input(f"Name #{i+1}", key=f"emp_name_{a}_{i}_{current_site}", placeholder="Full name")
            with ec2:
                emp_pay = st.number_input(f"Amount (KES) #{i+1}", min_value=0, step=100, key=f"emp_pay_{a}_{i}_{current_site}")
            employed_records.append({"name": emp_name, "amount": emp_pay})

    with p2:
        st.markdown("🔧 **Casual Workers**")
        num_casuals = st.number_input("Number of Casuals", min_value=0, max_value=20, step=1, key=f"num_cas_{a}_{current_site}")
        casual_records = []
        for i in range(int(num_casuals)):
            cc1, cc2 = st.columns(2)
            with cc1:
                cas_name = st.text_input(f"Name #{i+1}", key=f"cas_name_{a}_{i}_{current_site}", placeholder="Full name")
            with cc2:
                cas_pay = st.number_input(f"Amount (KES) #{i+1}", min_value=0, step=100, key=f"cas_pay_{a}_{i}_{current_site}")
            casual_records.append({"name": cas_name, "amount": cas_pay})

    employed_summary = "; ".join([f"{r['name']} (KES {r['amount']:,})" for r in employed_records if r['name']]) or "-"
    casuals_summary  = "; ".join([f"{r['name']} (KES {r['amount']:,})" for r in casual_records  if r['name']]) or "-"
    personnel_str = f"EMPLOYED: {employed_summary} | CASUALS: {casuals_summary}"

    area_remarks = st.text_area(
        "Remarks",
        key=f"remarks_{a}_{current_site}",
        placeholder="Challenges, causes, solutions, preventive measures"
    )

    areas_data.append({
        "name": area_name or f"Area {a+1}",
        "status": area_status,
        "improvements": area_improvements,
        "q1": q1, "q2": q2, "q3": q3, "q4": q4,
        "personnel": personnel_str,
        "remarks": area_remarks
    })

    st.divider()

# --- 4. ACTIONS ---
btn_col1, btn_col2, btn_col3 = st.columns(3)

# Build flat rows for CSV (one row per area)
rows = []
for area in areas_data:
    rows.append({
        "Report_Date": report_date,
        "Facility": current_site,
        "Area_of_Concentration": area["name"],
        "Current_Status": area["status"],
        "Improvements_Needed": area["improvements"],
        "Q1_Budget": area["q1"],
        "Q2_Budget": area["q2"],
        "Q3_Budget": area["q3"],
        "Q4_Budget": area["q4"],
        "Personnel": area["personnel"],
        "Remarks": area["remarks"],
        "Entry_Timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    })
df = pd.DataFrame(rows)

with btn_col1:
    if st.button(f"✅ Submit & Save {current_site} Report", use_container_width=True):
        file_path = "facility_master_log.csv"
        try:
            if not os.path.isfile(file_path):
                df.to_csv(file_path, index=False)
            else:
                df.to_csv(file_path, mode='a', header=False, index=False)
            st.balloons()
            st.success(f"Successfully logged {len(areas_data)} area(s) for {current_site}.")
        except Exception as e:
            st.error(f"Error saving data: {e}")

with btn_col2:
    single_csv = df.to_csv(index=False).encode("utf-8")
    st.download_button(
        label="⬇️ Download as CSV",
        data=single_csv,
        file_name=f"{current_site.replace(' ', '_')}_{report_date}.csv",
        mime="text/csv",
        use_container_width=True
    )

with btn_col3:
    word_buffer = generate_word_report(
        facility=current_site,
        report_date=report_date,
        areas=areas_data
    )
    st.download_button(
        label="⬇️ Download as Word (.docx)",
        data=word_buffer,
        file_name=f"{current_site.replace(' ', '_')}_{report_date}.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        use_container_width=True
    )

# --- 5. CREWAI NOTE ---
st.divider()
with st.expander("How to connect this to CrewAI"):
    st.write("""
    1. **Data Tool:** Give your CrewAI agents a `FileReadTool` pointing to `facility_master_log.csv`.
    2. **The Analyst Agent:** Filter by 'Facility' and 'Area_of_Concentration' columns.
    3. **The Manager Agent:** Aggregate budget and status columns for weekly/monthly/yearly reports.
    """)
